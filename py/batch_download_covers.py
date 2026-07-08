#!/usr/bin/env python3
"""
批量封面下载脚本
1. 从 merged.xlsx 读取所有游戏
2. 检测哪个游戏缺封面（文件不存在或使用 placeholder）
3. 从 retro_en_translated.json + 现有库存 JSON 中查找英文名
4. 从 libretro CDN 批量下载封面
5. 输出缺失报告

Usage:
    python py/batch_download_covers.py [--workers 10] [--timeout 15]
"""

import json
import os
import re
import ssl
import sys
import time
import urllib.parse
import urllib.request
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(ROOT, 'res', 'merged.xlsx')
COVERS_DIR = os.path.join(ROOT, 'res', 'covers')
TRANSLATED_PATH = os.path.join(ROOT, 'retro_en_translated.json')

LIBRETRO_HOST = "https://thumbnails.libretro.com"

PLATFORM_SYSTEM_MAP = {
    "FC/NES": "Nintendo - Nintendo Entertainment System",
    "SFC/SNES": "Nintendo - Super Nintendo Entertainment System",
    "GBA": "Nintendo - Game Boy Advance",
    "GBC/GB": "Nintendo - Game Boy Color",
    "N64": "Nintendo - Nintendo 64",
    "NDS": "Nintendo - Nintendo DS",
    "3DS": "Nintendo - Nintendo 3DS",
    "PS1": "Sony - PlayStation",
    "PS2": "Sony - PlayStation 2",
    "PS3": "Sony - PlayStation 3",
    "PSP": "Sony - PSP",
    "WII": "Nintendo - Wii",
    "NGC": "Nintendo - Nintendo GameCube",
    "MD": "Sega - Mega Drive - Genesis",
    "PCE": "NEC - PC Engine - TurboGrafx 16",
    "NGP": "SNK - Neo Geo Pocket Color",
    "WS": "Bandai - WonderSwan",
    "DOS": "DOS",
}

REGION_SUFFIXES = [
    "(USA)",
    "(USA, Europe)",
    "(USA) (En,Fr,De)",
    "(USA, Europe) (En,Fr,De)",
    "(Europe)",
    "(Europe) (En,Fr,De)",
    "(Europe) (En,Fr,De,Es,It,Nl)",
    "(Japan)",
    "(World)",
    "(World) (En,Ja)",
    "(USA) (En,Ja)",
    "(Japan) (En,Ja)",
]


def sanitize_filename(name):
    return re.sub(r'[<>:"/\\|?*]', "_", name).strip().rstrip(".") or "cover"


def build_libretro_urls(search_name, system):
    candidates = []
    base = search_name
    if base.lower().endswith(".png"):
        base = base[:-4]
    for suffix in ["", *REGION_SUFFIXES]:
        fname = f"{base} {suffix}.png" if suffix else f"{base}.png"
        url = f"{LIBRETRO_HOST}/{urllib.parse.quote(system, safe='')}/Named_Boxarts/{urllib.parse.quote(fname, safe='')}"
        candidates.append(url)
    return candidates


def download_url(url, output_path, timeout=15):
    context = ssl.create_default_context()
    context.check_hostname = False
    context.verify_mode = ssl.CERT_NONE
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    })
    try:
        with urllib.request.urlopen(req, context=context, timeout=timeout) as resp:
            data = resp.read()
        if len(data) < 2000:
            return False
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with open(output_path, "wb") as f:
            f.write(data)
        return True
    except Exception:
        return False


def load_excel_games():
    """Load all games from merged.xlsx using openpyxl."""
    from openpyxl import load_workbook
    wb = load_workbook(EXCEL_PATH, read_only=True)
    games = []
    for sname in wb.sheetnames:
        if sname == '主目录':
            continue
        ws = wb[sname]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 2:
                continue
            cover = str(row[0] or '').strip()
            name = str(row[1] or '').strip()
            plat = str(row[4] or '').strip() if len(row) > 4 else ''
            if not name:
                continue
            games.append({
                'name': name,
                'cover': cover,
                'plat': plat,
                'sheet': sname,
            })
    wb.close()
    return games


def load_en_mapping():
    """Build reverse mapping: Chinese name -> list of (English name, system)."""
    mapping = defaultdict(list)

    # Load from retro_en_translated.json
    if os.path.exists(TRANSLATED_PATH):
        with open(TRANSLATED_PATH, 'r', encoding='utf-8') as f:
            for entry in json.load(f):
                cn = entry.get('cn', '')
                en = entry.get('name', '')
                if cn and en:
                    mapping[cn].append((en, None))

    # Load from GBA inventory (has search + system fields)
    inventory_paths = [
        ('tmp/gba497_inventory.json', None),
        ('tmp/gbc_inventory.json', None),
    ]
    for inv_path, _ in inventory_paths:
        path = os.path.join(ROOT, inv_path)
        if os.path.exists(path):
            with open(path, 'r', encoding='utf-8') as f:
                for entry in json.load(f):
                    cn = entry.get('name', '')
                    search = entry.get('search', '')
                    system = entry.get('system', '')
                    if cn and search and search != cn:
                        mapping[cn].append((search, system))

    return mapping


def get_libretro_system(plat):
    for key, val in PLATFORM_SYSTEM_MAP.items():
        if key in plat or plat in key:
            return val
    return None


def download_one_game(game, en_mapping, stats):
    name = game['name']
    plat = game['plat']
    safe_name = sanitize_filename(name)
    out_path = os.path.join(COVERS_DIR, f"{safe_name}.png")

    # Skip if already exists and valid
    if os.path.exists(out_path) and os.path.getsize(out_path) > 2000:
        stats['existing'] += 1
        return True

    # Find English name candidates
    candidates = en_mapping.get(name, [])
    systems_tried = set()

    # Also try from direct match by name variants
    for search_name, system in candidates:
        libretro_sys = system or get_libretro_system(plat)
        if not libretro_sys:
            continue
        if libretro_sys in systems_tried:
            continue
        systems_tried.add(libretro_sys)

        urls = build_libretro_urls(search_name, libretro_sys)
        for url in urls:
            if download_url(url, out_path):
                stats['downloaded'] += 1
                return True
            time.sleep(0.1)

    # Try with just the Chinese name (CDN might have it)
    libretro_sys = get_libretro_system(plat)
    if libretro_sys and libretro_sys not in systems_tried:
        urls = build_libretro_urls(name, libretro_sys)
        for url in urls:
            if download_url(url, out_path):
                stats['downloaded'] += 1
                return True
            time.sleep(0.1)

    stats['missing'] += 1
    return False


def main():
    import argparse
    parser = argparse.ArgumentParser(description="批量下载游戏封面")
    parser.add_argument('--workers', type=int, default=10, help='并发下载数')
    parser.add_argument('--timeout', type=int, default=15, help='下载超时秒数')
    parser.add_argument('--limit', type=int, default=0, help='限制处理数量(调试用)')
    args = parser.parse_args()

    print("=" * 60)
    print("读取 merged.xlsx...")
    games = load_excel_games()
    print(f"共读取 {len(games)} 个游戏")

    # Filter only games with missing covers
    missing_games = []
    for g in games:
        cover = g['cover']
        if not cover or 'placeholder' in cover:
            missing_games.append(g)
        else:
            fname = os.path.basename(cover.replace('\\', '/'))
            fpath = os.path.join(COVERS_DIR, fname)
            if not os.path.exists(fpath) or os.path.getsize(fpath) <= 2000:
                missing_games.append(g)

    print(f"缺封面的游戏: {len(missing_games)}")

    if args.limit > 0:
        missing_games = missing_games[:args.limit]
        print(f"限制处理: {args.limit} 个")

    print("加载英文名映射...")
    en_mapping = load_en_mapping()
    print(f"映射表: {len(en_mapping)} 个中文名")

    print("=" * 60)
    print(f"开始下载 (并发 {args.workers})...")
    print("=" * 60)

    stats = {'downloaded': 0, 'existing': 0, 'missing': 0, 'total': len(missing_games)}
    failed_games = []

    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = {}
        for game in missing_games:
            fut = executor.submit(download_one_game, game, en_mapping, stats)
            futures[fut] = game

        for i, fut in enumerate(as_completed(futures), 1):
            game = futures[fut]
            try:
                ok = fut.result()
                if not ok:
                    failed_games.append(game)
            except Exception as e:
                failed_games.append(game)

            if i % 50 == 0 or i == len(futures):
                pct = i / len(futures) * 100
                print(f"  进度: {i}/{len(futures)} ({pct:.0f}%) | "
                      f"已下载: {stats['downloaded']} | "
                      f"已存在: {stats['existing']} | "
                      f"仍缺失: {stats['missing']}")

    print("=" * 60)
    print(f"完成! 下载: {stats['downloaded']}, "
          f"已存在: {stats['existing']}, "
          f"仍缺失: {stats['missing']}")

    # Save failed list
    if failed_games:
        report_path = os.path.join(ROOT, 'tmp', 'still_missing_covers.json')
        with open(report_path, 'w', encoding='utf-8') as f:
            json.dump(failed_games, f, ensure_ascii=False, indent=2)
        print(f"\n缺失报告: {report_path}")
        print(f"前20个缺失游戏:")
        for g in failed_games[:20]:
            print(f"  {g['name']} ({g['plat']})")


if __name__ == '__main__':
    main()
