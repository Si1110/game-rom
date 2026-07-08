#!/usr/bin/env python3
"""
更新 merged.xlsx 中封面路径：对已有封面文件但 Excel 里是 placeholder 的游戏，修正路径。
"""
import os, re
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(ROOT, 'res', 'merged.xlsx')
COVERS = os.path.join(ROOT, 'res', 'covers')

def sanitize(name):
    return re.sub(r'[\\/:*?"<>|]', '_', name)

print("Loading Excel...")
wb = load_workbook(EXCEL)
updated = 0
total = 0

for sname in wb.sheetnames:
    if sname == '主目录':
        continue
    ws = wb[sname]
    for row in ws.iter_rows(min_row=2):
        cells = [c.value for c in row]
        if not cells or len(cells) < 2:
            continue
        cover = str(cells[0] or '').strip() if cells[0] is not None else ''
        name = str(cells[1] or '').strip() if cells[1] is not None else ''
        if not name:
            continue
        total += 1

        if not cover or 'placeholder' in cover:
            safe = sanitize(name)
            expected = f'{safe}.png'
            fpath = os.path.join(COVERS, expected)
            if os.path.exists(fpath) and os.path.getsize(fpath) > 2000:
                new_cover = f'../res/covers/{expected}'
                row[0].value = new_cover
                updated += 1
                if updated <= 5 or updated % 1000 == 0:
                    print(f"  [{sname}] {name} -> {new_cover}")

print(f"\nTotal games: {total}")
print(f"Cover paths updated: {updated}")

if updated > 0:
    wb.save(EXCEL)
    print(f"Saved: {EXCEL}")
else:
    print("No changes needed.")
